"""spec 2026-08-11-library-attachment-rag：附件正文解析工具路由（mock storage，不连 MinIO）。"""
import io
import json

import pytest
from fastapi.responses import JSONResponse

from agent.routes import parse_text as mod
from agent.routes.parse_text import ParseTextBody, parse_text


def _docx(*paras: str) -> bytes:
    from docx import Document
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf_with_blank_page() -> bytes:
    """一页有字 + 一页空白（对文本提取而言，扫描图片页与空白页是同一回事）。"""
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(0, 10, "Tender body text that is clearly longer than the threshold")
    pdf.add_page()
    return bytes(pdf.output())


def _text_pdf(pages: int = 2) -> bytes:
    """整份都是文字层的 PDF（每页都远超可见字数门槛）→ 一页扫描页都没有。"""
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_font("helvetica", size=12)
    for i in range(pages):
        pdf.add_page()
        pdf.cell(0, 10, f"Page {i + 1}: tender body text that is clearly longer than threshold")
    return bytes(pdf.output())


def _blank_pdf(pages: int) -> bytes:
    """整份提不出文字的 PDF（等价扫描件：对文本提取而言空白页与扫描图片页是同一回事）。"""
    from fpdf import FPDF
    pdf = FPDF()
    for _ in range(pages):
        pdf.add_page()
    return bytes(pdf.output())


_RECOGNIZED = "识别文字·统一社会信用代码 91110000MA01ABCD2X·法定代表人授权委托书盖章页"


@pytest.fixture
def ocr_stub(monkeypatch):
    """OCR 容器的 HTTP 桩 + 配上地址。换的是 **ocr 模块看到的那个 httpx**（不是 httpx 本身）：
    往真 httpx 上 setattr 等于改全局，测试期间任何别的组件构造 AsyncClient 都会拿到这份桩
    （手法与 tests/parsing/conftest.py 的 ocr_stub 同源）。"""
    import types

    import httpx

    from agent.parsing import ocr as ocr_mod

    class _Stub:
        def __init__(self):
            self.requests: list[dict] = []
            self.clients = 0

        def __call__(self, request: httpx.Request) -> httpx.Response:
            assert request.url.path == "/ocr"
            self.requests.append(json.loads(request.content))
            return httpx.Response(200, json={"text": _RECOGNIZED})

    stub = _Stub()
    real_client = httpx.AsyncClient
    monkeypatch.setattr(ocr_mod.settings, "ocr_base_url", "http://ocr.test:8100/")

    def _client(**kw):
        stub.clients += 1
        return real_client(transport=httpx.MockTransport(stub), **kw)

    monkeypatch.setattr(ocr_mod, "httpx", types.SimpleNamespace(AsyncClient=_client))
    return stub


@pytest.fixture
def fake_read(monkeypatch):
    """把 read_bytes 换成内存桩，返回它记下的 key。"""
    calls: list[str] = []
    data = {"bytes": b""}

    def _read(key: str) -> bytes:
        calls.append(key)
        return data["bytes"]

    monkeypatch.setattr(mod, "read_bytes", _read)
    return type("F", (), {"calls": calls, "set": staticmethod(lambda b: data.update(bytes=b))})


async def test_docx_returns_plain_text(fake_read):
    fake_read.set(_docx("零信任统一身份认证技术方案", "采用持续验证与最小权限原则"))
    resp = await parse_text(ParseTextBody(key="uploads/u1/f1/零信任.docx"))
    assert resp["kind"] == "docx" and resp["no_text"] is False
    assert "零信任统一身份认证技术方案" in resp["text"]
    assert "最小权限原则" in resp["text"]
    assert resp["chars"] == len(resp["text"]) and resp["truncated"] is False


async def test_truncates_at_max_chars(fake_read):
    fake_read.set(_docx("零" * 5000))
    resp = await parse_text(ParseTextBody(key="a.docx", max_chars=100))
    assert len(resp["text"]) == 100 and resp["truncated"] is True and resp["chars"] == 100


async def test_max_chars_is_capped(fake_read):
    """调用方传天文数字也不会让一份大文件整篇过网（硬顶在路由侧）。"""
    fake_read.set(_docx("零" * 5000))
    resp = await parse_text(ParseTextBody(key="a.docx", max_chars=10_000_000))
    assert len(resp["text"]) <= mod._MAX_CHARS_CAP


async def test_pdf_reports_image_pages(fake_read):
    fake_read.set(_pdf_with_blank_page())
    resp = await parse_text(ParseTextBody(key="b.pdf"))
    assert resp["kind"] == "pdf" and resp["image_pages"] == 1
    assert "Tender body text" in resp["text"] and resp["no_text"] is False


async def test_scanned_pdf_reports_no_text_when_ocr_unconfigured(fake_read):
    """OCR 未配置（这套环境没部署识别服务）→ 整段跳过，如实报 no_text，不失败。"""
    fake_read.set(_blank_pdf(1))
    resp = await parse_text(ParseTextBody(key="scan.pdf"))
    assert resp["no_text"] is True and resp["text"] == "" and resp["chars"] == 0
    assert resp["ocr_pages"] == 0 and resp["image_pages"] == 1


async def test_unsupported_extension_is_422(fake_read):
    fake_read.set(b"whatever")
    resp = await parse_text(ParseTextBody(key="note.txt"))
    assert isinstance(resp, JSONResponse) and resp.status_code == 422
    assert b"unsupported" in resp.body


async def test_broken_bytes_is_422(fake_read):
    """坏字节不裸崩：内容与扩展名不符先被魔数校验拦成 unsupported，其余异常兜成 parse_failed。"""
    fake_read.set(b"not a real docx")
    resp = await parse_text(ParseTextBody(key="broken.docx"))
    assert isinstance(resp, JSONResponse) and resp.status_code == 422
    assert b"unsupported" in resp.body or b"parse_failed" in resp.body


async def test_storage_failure_is_422(fake_read, monkeypatch):
    """取件本身失败（对象没了/MinIO 抽风）同样是 422，不是 500——对调用方是「没正文」。"""
    def _boom(key: str) -> bytes:
        raise RuntimeError("NoSuchKey")

    monkeypatch.setattr(mod, "read_bytes", _boom)
    resp = await parse_text(ParseTextBody(key="gone.docx"))
    assert isinstance(resp, JSONResponse) and resp.status_code == 422
    assert b"parse_failed" in resp.body


async def test_returned_text_carries_no_system_note(fake_read):
    """返回文本里绝不含系统注记：它会进 RAG 索引被当参考资料喂给写手，抄进标书=审查报残留。"""
    from agent.parsing.types import SYSTEM_NOTE_PREFIX
    fake_read.set(_docx("零" * 5000))
    resp = await parse_text(ParseTextBody(key="a.docx", max_chars=100))
    assert SYSTEM_NOTE_PREFIX not in resp["text"]


@pytest.fixture
def ocr_reads(monkeypatch):
    """ocr_scanned_pages 进门要把字节重新取一次（它有自己那份 read_bytes 引用）。"""
    from agent.parsing import ocr as ocr_mod
    data = {"bytes": b""}
    monkeypatch.setattr(ocr_mod, "read_bytes", lambda key: data["bytes"])
    return data


async def test_scanned_pdf_is_ocred_into_text(fake_read, ocr_reads, ocr_stub):
    """①扫描版 PDF → 识别文字进 text（App 据此写回附件并入索引）。"""
    pdf = _blank_pdf(2)
    fake_read.set(pdf)
    ocr_reads["bytes"] = pdf

    resp = await parse_text(ParseTextBody(key="扫描版资质.pdf"))

    assert len(ocr_stub.requests) == 2
    assert resp["no_text"] is False and _RECOGNIZED in resp["text"]
    assert resp["ocr_pages"] == 2 and resp["image_pages"] == 0


async def test_text_pdf_never_calls_ocr(fake_read, ocr_reads, ocr_stub):
    """②文字版 PDF 一次 OCR 都不该发（连 HTTP 客户端都不该构造）。"""
    pdf = _text_pdf(3)
    fake_read.set(pdf)
    ocr_reads["bytes"] = pdf

    resp = await parse_text(ParseTextBody(key="文字版.pdf"))

    assert ocr_stub.requests == [] and ocr_stub.clients == 0
    assert resp["no_text"] is False and resp["ocr_pages"] == 0


async def test_ocr_can_be_switched_off_per_call(fake_read, ocr_reads, ocr_stub):
    """ocr=false（调用方明说不要识别）→ 零调用，退回只解析文字层。"""
    pdf = _blank_pdf(1)
    fake_read.set(pdf)
    ocr_reads["bytes"] = pdf

    resp = await parse_text(ParseTextBody(key="scan.pdf", ocr=False))

    assert ocr_stub.requests == [] and resp["no_text"] is True


async def test_ocr_text_carries_no_system_note(fake_read, ocr_reads, ocr_stub):
    """识别拼回时的页首注记（【系统注记·扫描页识别 第N页】）绝不能跟着进索引。"""
    from agent.parsing.types import SYSTEM_NOTE_PREFIX
    pdf = _blank_pdf(1)
    fake_read.set(pdf)
    ocr_reads["bytes"] = pdf

    resp = await parse_text(ParseTextBody(key="scan.pdf"))

    assert _RECOGNIZED in resp["text"] and SYSTEM_NOTE_PREFIX not in resp["text"]


async def test_budget_exhausted_before_start_falls_back_to_text_layer(
        fake_read, ocr_reads, ocr_stub, monkeypatch):
    """④时长预算在排队里就耗光 → 只入文字层，不报错（下次保存重试）。"""
    pdf = _blank_pdf(1)
    fake_read.set(pdf)
    ocr_reads["bytes"] = pdf
    monkeypatch.setattr(mod, "_OCR_BUDGET_S", 0)   # deadline 进门即已过期

    resp = await parse_text(ParseTextBody(key="scan.pdf"))

    assert ocr_stub.requests == [] and resp["no_text"] is True   # 200，不是错误


async def test_deadline_mid_stream_keeps_partial_pages(
        fake_read, ocr_reads, ocr_stub, monkeypatch):
    """④时长预算在识别中途到点 → 已识别的页照常入索引，剩下的记为仍看不见，**绝不整条失败**。

    用假时钟驱动（并发 2 路 = 每块 2 页）：第一块开始时还在预算内，第二块开始时已过期。"""
    import types

    from agent.parsing import ocr as ocr_mod
    pdf = _blank_pdf(4)
    fake_read.set(pdf)
    ocr_reads["bytes"] = pdf
    ticks = iter([0.0, 0.0, 10 ** 9] + [10 ** 9] * 50)  # 进门检查、第一块、第二块…
    monkeypatch.setattr(ocr_mod, "time", types.SimpleNamespace(monotonic=lambda: next(ticks)))

    resp = await parse_text(ParseTextBody(key="扫描件.pdf"))

    assert len(ocr_stub.requests) == 2                  # 只做完第一块就收手
    assert resp["ocr_pages"] == 2 and resp["image_pages"] == 2
    assert _RECOGNIZED in resp["text"] and resp["no_text"] is False


def test_router_is_mounted():
    """接线必须是真的（本项目「写了但没接上」翻过多次车）。"""
    from agent.app import create_app
    app = create_app()
    paths = set()
    for route in app.routes:
        if type(route).__name__ == '_IncludedRouter' and hasattr(route, 'original_router'):
            for r in route.original_router.routes:
                if hasattr(r, 'path'):
                    paths.add(r.path)
    assert "/tools/parse-text" in paths

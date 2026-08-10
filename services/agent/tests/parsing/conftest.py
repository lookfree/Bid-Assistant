import io
import json

import httpx
import pytest


def recognized_text(n: int) -> str:
    """桩的识别结果。必须**够读**（≥ parsers._MIN_VISIBLE_CHARS），否则那一页/那张图照旧算
    看不见——真实证照认出来的是整版文字，一两个字的才是异常路径（另有用例专测）。"""
    return f"识别文字{n}·法定代表人身份证明及授权委托书盖章页"


class OcrStub:
    """OCR 容器的 HTTP 桩：记录每次请求体，按到达序给答复（reply 可换成失败/超时）。

    扫描页与 docx 内嵌图两条链路共用同一个 httpx 客户端与同一套韧性口径，桩也只该有一份。"""

    def __init__(self):
        self.requests: list[dict] = []
        self.reads: list[str] = []      # 取字节的 key，供「该跳过时连字节都不该取」这类断言
        self.reply = lambda n, body: httpx.Response(200, json={"text": recognized_text(n)})

    def __call__(self, request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/ocr"
        body = json.loads(request.content)
        self.requests.append(body)
        return self.reply(len(self.requests), body)


@pytest.fixture
def ocr_stub(monkeypatch):
    """配上 OCR 地址 + 把传输层换成桩（仍真走 httpx 的请求/响应/异常路径）。"""
    from agent.parsing import ocr as ocr_mod

    stub = OcrStub()
    real_client = httpx.AsyncClient
    monkeypatch.setattr(ocr_mod.settings, "ocr_base_url", "http://ocr.test:8100/")
    monkeypatch.setattr(ocr_mod.httpx, "AsyncClient",
                        lambda **kw: real_client(transport=httpx.MockTransport(stub), **kw))
    return stub


@pytest.fixture
def docgen():
    """三类型测试文件生成器（docx/xlsx/pdf），供 parsing 测试共用。"""
    class _Gen:
        @staticmethod
        def docx(*paras: str) -> bytes:
            from docx import Document
            d = Document()
            for p in paras:
                d.add_paragraph(p)
            buf = io.BytesIO()
            d.save(buf)
            return buf.getvalue()

        @staticmethod
        def xlsx() -> bytes:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.append(["评分项", "分值"])
            ws.append(["技术标", 60])
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()

        @staticmethod
        def pdf(text: str) -> bytes:
            from fpdf import FPDF
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("helvetica", size=12)
            pdf.cell(0, 10, text)
            return bytes(pdf.output())

    return _Gen
